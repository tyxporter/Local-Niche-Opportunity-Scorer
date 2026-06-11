"""Branded .docx export (§4.5) — the firm-facing deliverable.

Renders a `Brief` (the 6-section content model from brief.py) into a Carson
Wealth–branded Word document, 1–2 pages, named
    [Firm]-Local-Opportunity-Brief-[YYYY-MM-DD].docx
that stands alone as a firm-facing artifact.

Boundaries this module HOLDS:
  * It renders ONLY the Brief's six sections. It never imports or includes the
    delivery wrapper (delivery.py) — that is app-only and never in the .docx
    (§4.5). Separation is structural: no shared state, no cross-import.
  * The Brief model carries no scorer numbers or labels (D1/D2), so none can
    reach the page.

Brand tokens come from brand.py (§5). The Carson Wealth logo is a flagged
placeholder until Ty supplies it: if brand.LOGO_PATH is unset we render a
visible placeholder rather than silently shipping an unbranded header.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError as exc:  # pragma: no cover
    raise ImportError("python-docx is required for .docx export. "
                      "`pip install python-docx`.") from exc

from . import brand
from .brief import Brief

DELIVERABLE_TITLE = "Local Marketing Opportunity Brief"

_NAVY = RGBColor.from_string(brand.NAVY.lstrip("#"))
_GOLD = RGBColor.from_string(brand.GOLD.lstrip("#"))
_GRAY = RGBColor.from_string(brand.GRAY_DARK.lstrip("#"))
_BODY = RGBColor.from_string("222b30")


# --- low-level brand helpers ------------------------------------------------
def _set_font(run, *, size, color, bold=False, name=brand.FONT_PRIMARY):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Pin the font across ascii/hAnsi/cs slots (fallback chain documented in §5).
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(slot), name)


def _gold_rule(paragraph):
    """Gold bottom border under a paragraph — the chevron/diagonal accent (§5),
    expressed as a clean horizontal rule (on-brand: angular/structural)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), brand.GOLD.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


def _heading(doc, text, *, size=13):
    p = doc.add_paragraph()
    p.space_after = Pt(2)
    run = p.add_run(text)
    _set_font(run, size=size, color=_NAVY, bold=True)
    _gold_rule(p)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def _body_paragraph(doc, text, *, size=10.5, color=_BODY, bullet=False, italic=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    _set_font(run, size=size, color=color)
    run.font.italic = italic
    return p


# --- header -----------------------------------------------------------------
def _render_header(doc, brief: Brief):
    logo = brand.LOGO_PATH
    if logo and Path(logo).exists():
        doc.add_picture(str(logo), width=Inches(1.8))
    else:
        ph = doc.add_paragraph()
        r = ph.add_run("[ Carson Wealth logo — pending Ty ]")
        _set_font(r, size=9, color=_GRAY)
        r.font.italic = True

    title = doc.add_paragraph()
    tr = title.add_run(DELIVERABLE_TITLE)
    _set_font(tr, size=20, color=_NAVY, bold=True)
    title.paragraph_format.space_after = Pt(2)

    sub = doc.add_paragraph()
    sr = sub.add_run(f"{brief.firm_legal_name}")
    _set_font(sr, size=13, color=_GOLD, bold=True)
    sub.paragraph_format.space_after = Pt(0)

    dt = doc.add_paragraph()
    dr = dt.add_run(f"Prepared {brief.generated_on}")
    _set_font(dr, size=9, color=_GRAY)
    dt.paragraph_format.space_after = Pt(6)
    _gold_rule(dt)


# --- public API -------------------------------------------------------------
def render_brief_docx(brief: Brief, *, out_dir: str | Path = ".") -> Path:
    """Render the Brief to a branded .docx and return the written path."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Base style (body: light weight, generous leading — §5).
    normal = doc.styles["Normal"]
    normal.font.name = brand.FONT_PRIMARY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _BODY

    _render_header(doc, brief)

    for sec in brief.sections:
        _heading(doc, f"{sec.number}. {sec.title}")
        is_snapshot = sec.number == 1
        is_footer = sec.number == 6
        for line in (sec.body.split("\n") if sec.body else []):
            line = line.rstrip()
            if not line:
                continue
            if is_footer:
                _body_paragraph(doc, line, size=8.5, color=_GRAY, italic=True)
            elif is_snapshot and not line.lower().startswith("note:"):
                _body_paragraph(doc, line, bullet=True)
            else:
                _body_paragraph(doc, line)

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{brief.filename_base()}.docx"
    doc.save(str(path))
    return path
