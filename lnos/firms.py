"""Firm spine (§4.0, Phase 1).

Loads the canonical firm list from the verified Carson Wealth AUM workbook —
the ONLY permitted source for the firm spine (do not pull from any other
export, per §1 / D6).

What the workbook actually contains (verified by reading it):
  * Sheet "Firm AUM (Verified)": 45 firms with verified advisory AUM,
    total $26,702,566,296.72.
  * Sheet "Exceptions — Need Confirmation":
      - 2 UNMATCHED firms (Johnson City, Las Vegas) — AUM pending. We leave
        AUM null and NEVER infer it (§4.0, §8).
      - 2 HOUSE accounts (Carson Group, Corporate Accounts) — excluded from
        the firm set by design.

=> The scoreable firm spine is 45 + 2 = 47 firms. This count is derived from
   the file, never hardcoded. (Brief v2.3 locks the portfolio at "47 wholly-
   owned advisor firms (49 rows − 2 house accounts)" — this matches the loader
   exactly; the earlier "49" framing counted the two excluded house accounts.)

Geography (city/state/zip/county/CBSA) is NOT in the AUM workbook. Per brief
v2.3 it is a SEPARATE spine input, joined from the canonical companion file
Carson-Firm-Geography-Template-2026-06-11.xlsx (Ty completes from Salesforce).
We do NOT infer geography from firm names. The template ships with County/CBSA
blank, so firms load with geography unresolved and flagged; the brief/FRED
layers fail loud when geography is required (county + CBSA must be specific
enough for FRED series). The template's "City (suggested — VERIFY)" column is
UNCONFIRMED and is carried as `city_suggested` only — never promoted to the
authoritative `city` until the row's County + CBSA are filled.

AUM is descriptive INPUT data only — never a scorer signal, never blocks
scoring (a null AUM is fine).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - openpyxl is a core dep
    raise ImportError(
        "openpyxl is required to load the firm spine. `pip install openpyxl`."
    ) from exc


# --- Canonical file locations -----------------------------------------------
_DATA_DIR = Path(__file__).resolve().parent.parent / "data"
SPINE_XLSX = _DATA_DIR / "Carson-WhollyOwned-Firm-AUM-Verified-2026-06-10.xlsx"
GEOGRAPHY_XLSX = _DATA_DIR / "Carson-Firm-Geography-Template-2026-06-11.xlsx"
GEOGRAPHY_SHEET = "Firm Geography (to complete)"

# Verified expectations (used by the loader's self-check, not as a source of
# methodology). If the file ever stops matching these, we fail loud.
EXPECTED_MATCHED_COUNT = 45
EXPECTED_MATCHED_TOTAL_USD = 26702566296.72


@dataclass
class Firm:
    """A wholly-owned advisor firm in the spine."""

    firm_id: str                       # slug derived from roster name
    roster_name: str                   # canonical name (workbook col B)
    aum_line_name: Optional[str] = None  # AUM-report close name (col C)
    aum_usd: Optional[float] = None    # advisory AUM (None if pending)
    aum_pending: bool = False          # True -> Johnson City / Las Vegas
    is_house_account: bool = False     # True -> excluded from the firm set

    # Geography — joined from the canonical geography template (brief v2.3);
    # None until Ty completes it from Salesforce.
    city: Optional[str] = None         # authoritative (only set once row complete)
    state: Optional[str] = None
    zip: Optional[str] = None
    county: Optional[str] = None
    cbsa: Optional[str] = None         # metro CBSA (name or code)
    cbsa_code: Optional[str] = None    # numeric CBSA code if provided
    city_suggested: Optional[str] = None  # UNCONFIRMED hint from office name

    notes: str = ""

    @property
    def has_geography(self) -> bool:
        """True when geography is specific enough for FRED county/CBSA series."""
        return bool(self.county and (self.cbsa or self.cbsa_code))

    @property
    def aum_usd_m(self) -> Optional[float]:
        return None if self.aum_usd is None else round(self.aum_usd / 1_000_000, 2)


def _slug(name: str) -> str:
    s = name.strip().lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


# --- Workbook parsing -------------------------------------------------------
def _find_header_row(rows, required) -> int:
    """Locate the header row containing all `required` cell labels."""
    req = {r.lower() for r in required}
    for i, row in enumerate(rows):
        cells = {str(c).strip().lower() for c in row if c is not None}
        if req.issubset(cells):
            return i
    raise ValueError(f"Could not find header row containing {required}")


def _parse_verified_sheet(ws) -> list[Firm]:
    rows = list(ws.iter_rows(values_only=True))
    hdr = _find_header_row(rows, {"Firm (roster name)", "AUM ($)"})
    header = [str(c).strip() if c is not None else "" for c in rows[hdr]]
    idx = {name: header.index(name) for name in header if name}
    firms: list[Firm] = []
    for row in rows[hdr + 1:]:
        roster = row[idx["Firm (roster name)"]]
        if roster is None:
            continue
        roster = str(roster).strip()
        if roster.upper().startswith("TOTAL"):  # summary row -> stop
            break
        aum = row[idx["AUM ($)"]] if idx.get("AUM ($)") is not None else None
        line_name = row[idx["Tied to AUM line"]] if "Tied to AUM line" in idx else None
        firms.append(
            Firm(
                firm_id=_slug(roster),
                roster_name=roster,
                aum_line_name=str(line_name).strip() if line_name else None,
                aum_usd=float(aum) if isinstance(aum, (int, float)) else None,
                aum_pending=False,
            )
        )
    return firms


def _parse_exceptions_sheet(ws):
    """Return (pending_firms, house_accounts) from the exceptions sheet."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = _find_header_row(rows, {"Firm", "Type"})
    header = [str(c).strip() if c is not None else "" for c in rows[hdr]]
    idx = {name: header.index(name) for name in header if name}
    pending: list[Firm] = []
    house: list[Firm] = []
    for row in rows[hdr + 1:]:
        name = row[idx["Firm"]]
        if name is None:
            continue
        name = str(name).strip()
        if name.lower().startswith("reminder"):  # footnote row
            continue
        ftype = str(row[idx["Type"]]).strip().upper() if row[idx["Type"]] else ""
        note_col = idx.get("Reason / Note")
        note = str(row[note_col]).strip() if note_col is not None and row[note_col] else ""
        if ftype == "UNMATCHED":
            pending.append(
                Firm(firm_id=_slug(name), roster_name=name,
                     aum_usd=None, aum_pending=True, notes=note)
            )
        elif ftype == "HOUSE":
            house.append(
                Firm(firm_id=_slug(name), roster_name=name,
                     is_house_account=True, notes=note)
            )
    return pending, house


# --- Geography join (canonical companion file, brief v2.3) ------------------
def _clean(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _load_geography() -> dict[str, dict]:
    """Read the geography template keyed by roster name. Missing file -> {}.

    Reads the authoritative columns (State / ZIP / County / CBHA-Metro) plus the
    UNCONFIRMED "City (suggested — VERIFY)" hint. Keyed by the exact roster name
    so it joins 1:1 with the AUM spine.
    """
    if not GEOGRAPHY_XLSX.exists():
        return {}
    wb = openpyxl.load_workbook(GEOGRAPHY_XLSX, data_only=True)
    ws = wb[GEOGRAPHY_SHEET] if GEOGRAPHY_SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = _find_header_row(rows, {"Firm (roster name)", "County"})
    header = [str(c).strip() if c is not None else "" for c in rows[hdr]]
    idx = {name: header.index(name) for name in header if name}
    name_col = idx["Firm (roster name)"]
    city_col = next((idx[h] for h in header if h.lower().startswith("city")), None)
    out: dict[str, dict] = {}
    for row in rows[hdr + 1:]:
        roster = row[name_col] if name_col < len(row) else None
        if roster is None:
            continue
        roster = str(roster).strip()
        if not roster or roster.lower().startswith(("legend", "•")):
            break
        out[roster] = {
            "city_suggested": _clean(row[city_col]) if city_col is not None else None,
            "state": _clean(row[idx["State"]]) if "State" in idx else None,
            "zip": _clean(row[idx["ZIP"]]) if "ZIP" in idx else None,
            "county": _clean(row[idx["County"]]) if "County" in idx else None,
            "cbsa": _clean(row[idx.get("CBSA / Metro", idx.get("CBSA"))])
            if ("CBSA / Metro" in idx or "CBSA" in idx) else None,
        }
    return out


def _apply_geography(firm: Firm, geo: dict) -> None:
    if not geo:
        return
    firm.state = geo.get("state")
    firm.zip = geo.get("zip")
    firm.county = geo.get("county")
    firm.cbsa = geo.get("cbsa")
    firm.city_suggested = geo.get("city_suggested")
    # Promote the suggested city to authoritative ONLY when the row is actually
    # complete (county + CBSA present). Until then we never treat the unverified
    # office-name hint as confirmed geography.
    if firm.has_geography:
        firm.city = geo.get("city_suggested")


# --- Public API -------------------------------------------------------------
def load_firms(
    *,
    include_pending: bool = True,
    include_house: bool = False,
    spine_path: Optional[Path] = None,
    verify: bool = True,
) -> list[Firm]:
    """Load the firm spine from the verified workbook + geography merge.

    By default returns the scoreable firm set (matched + pending, house
    accounts EXCLUDED). Set include_house=True only for reconciliation.
    """
    path = Path(spine_path) if spine_path else SPINE_XLSX
    if not path.exists():
        raise FileNotFoundError(
            f"Firm spine workbook not found at {path}. This is the canonical "
            f"firm source (§1) — it must be present. Do not substitute another export."
        )
    wb = openpyxl.load_workbook(path, data_only=True)
    matched = _parse_verified_sheet(wb["Firm AUM (Verified)"])
    pending, house = _parse_exceptions_sheet(wb["Exceptions — Need Confirmation"])

    if verify:
        _verify_spine(matched)

    geo = _load_geography()
    firms: list[Firm] = list(matched)
    if include_pending:
        firms += pending
    if include_house:
        firms += house
    for f in firms:
        _apply_geography(f, geo.get(f.roster_name, {}))
    return firms


def _verify_spine(matched: list[Firm]) -> None:
    """Fail loud if the workbook no longer matches the verified baseline."""
    n = len(matched)
    if n != EXPECTED_MATCHED_COUNT:
        raise ValueError(
            f"Firm spine self-check failed: expected {EXPECTED_MATCHED_COUNT} "
            f"matched firms, found {n}. The canonical workbook changed — stop "
            f"and reconcile before scoring (do not guess)."
        )
    total = sum(f.aum_usd for f in matched if f.aum_usd is not None)
    if abs(total - EXPECTED_MATCHED_TOTAL_USD) > 1.0:
        raise ValueError(
            f"Firm spine self-check failed: matched AUM total {total:.2f} != "
            f"expected {EXPECTED_MATCHED_TOTAL_USD:.2f}. Reconcile before use."
        )


def firms_missing_geography(firms: Optional[list[Firm]] = None) -> list[Firm]:
    """Firms that cannot yet support FRED county/CBSA series (open item F2)."""
    firms = firms if firms is not None else load_firms()
    return [f for f in firms if not f.has_geography]


def spine_summary() -> dict:
    """Counts/totals for reporting — purely descriptive."""
    all_firms = load_firms(include_pending=True, include_house=True)
    scoreable = [f for f in all_firms if not f.is_house_account]
    matched = [f for f in scoreable if not f.aum_pending]
    pending = [f for f in scoreable if f.aum_pending]
    house = [f for f in all_firms if f.is_house_account]
    return {
        "scoreable_firms": len(scoreable),
        "matched_with_aum": len(matched),
        "aum_pending": [f.roster_name for f in pending],
        "house_accounts_excluded": [f.roster_name for f in house],
        "matched_aum_total_usd": round(sum(f.aum_usd for f in matched), 2),
        "firms_missing_geography": len(firms_missing_geography(scoreable)),
    }


if __name__ == "__main__":  # quick manual check
    import json
    print(json.dumps(spine_summary(), indent=2))
